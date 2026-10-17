# -*- coding: utf-8 -*-
"""Les deroules de seance destines au responsable (Word).

   Registre administratif : pas de code, pas de jargon, pas d'emoji.
   Le responsable doit pouvoir lire le document sans rien connaitre
   a la programmation, et le transferer tel quel.
"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from deroules_donnees import SEANCES, RACINE

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

MARINE = RGBColor(0x1A, 0x1A, 0x2E)
VIOLET = RGBColor(0x6B, 0x2D, 0x8F)
GRIS = RGBColor(0x66, 0x66, 0x70)


def ombrer(cell, hexa):
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), hexa)
    cell._tc.get_or_add_tcPr().append(sh)


def construire(s):
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(1.8)
        sec.left_margin = sec.right_margin = Cm(2.0)

    def para(txt="", size=10.5, bold=False, color=None, after=6, italic=False, align=None):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after)
        if align is not None: p.alignment = align
        r = p.add_run(txt); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
        if color is not None: r.font.color.rgb = color
        return p

    def riche(morceaux, size=10.5, after=6, puce=False):
        p = doc.add_paragraph(style="List Bullet" if puce else None)
        p.paragraph_format.space_after = Pt(after)
        for txt, gras in morceaux:
            r = p.add_run(txt); r.bold = gras; r.font.size = Pt(size)
        return p

    # ------------------------------------------------------------ en-tete
    para("ACPROKIDS CODING CAMP · VISION PLÉNITUDES VIE", 8.5, True, VIOLET, after=2)
    para("PROGRAMME DE FORMATION — CODAGE PYTHON POUR LES JEUNES", 9, True, GRIS, after=8)
    para("Séance %d — « %s »" % (s["num"], s["titre"]), 17, True, MARINE, after=4)
    riche([("Date : ", True), ("%s, de 12 h à 14 h" % s["date"], False),
           ("        Encadrement : ", True), ("un enseignant par classe", False)], 9.5, after=2)
    riche([("Public : ", True),
           ("Jerusalem Geeks (9 – 12 ans) et Jeremiah Geeks (12 – 18 ans)", False)], 9.5, after=12)

    # ------------------------------------------------------------ objet
    riche([("Objet de la séance. ", True), (s["objet"], False)], after=12)

    # ------------------------------------------------------------ tableau
    para("LE PROGRAMME EN UN COUP D'ŒIL", 11, True, MARINE, after=6)
    t = doc.add_table(rows=1 + len(s["horaires"]), cols=3)
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, titre in enumerate(("Horaire", "Séquence", "Durée")):
        c = t.cell(0, i); c.text = ""
        r = c.paragraphs[0].add_run(titre); r.bold = True; r.font.size = Pt(9)
        ombrer(c, "EFEFF6")
    for li, (h, quoi, duree) in enumerate(s["horaires"], start=1):
        for ci, v in enumerate((h, quoi, duree)):
            c = t.cell(li, ci); c.text = ""
            r = c.paragraphs[0].add_run(v); r.font.size = Pt(9)
    t.columns[0].width = Cm(3.3); t.columns[1].width = Cm(10.6); t.columns[2].width = Cm(2.1)
    para("", 6, after=10)

    # ------------------------------------------------------------ deroule
    para("LE DÉROULÉ EN DÉTAIL", 11, True, MARINE, after=6)
    for titre, paragraphes in s["deroule"]:
        para(titre, 10.5, True, VIOLET, after=3)
        for p in paragraphes:
            para(p, 10.5, after=4)
        para("", 4, after=2)

    # ------------------------------------------------------------ travail
    para("TRAVAIL DONNÉ POUR LA SEMAINE", 11, True, MARINE, after=6)
    para(s["travail"], after=12)

    # ------------------------------------------------------------ acquis
    para(s["acquis"], after=4)

    sortie = os.path.join(RACINE, s["dossier"], "a-envoyer",
                          "deroule-seance-%d.docx" % s["num"])
    doc.save(sortie)
    return sortie


if __name__ == "__main__":
    for s in SEANCES:
        print("écrit :", construire(s))
