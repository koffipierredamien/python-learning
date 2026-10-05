# -*- coding: utf-8 -*-
"""Repart du classeur de l'utilisateur (le plus recent) et y rejoue nos ajouts :
   - l'onglet Revisions maison passe a un devoir par semaine, et il est rempli ;
   - les quatre retards excuses sont inscrits dans Observations.
"""
import openpyxl, datetime
from openpyxl.utils import get_column_letter

SRC = '/root/.claude/uploads/731be2d5-98da-5705-b521-79ab616d2443/98803b27-AcProKidsCodingCampSept2026.xlsx'
DST = '3-suivi/AcProKidsCodingCampSept2026.xlsx'
wb = openpyxl.load_workbook(SRC)

# =====================================================================
#  1. REVISIONS MAISON : un seul devoir par semaine
# =====================================================================
ws = wb['Révisions maison']
if ws.cell(3, 5).value == 'S01-2':                     # pas encore restructure
    for plage in [str(p) for p in ws.merged_cells.ranges if p.min_row in (1, 2)]:
        ws.unmerge_cells(plage)
    for c in range(47, 4, -2):
        ws.delete_cols(c)
    for i in range(22):
        ws.cell(3, 4 + i).value = 'S%02d' % (i + 1)
    TOT = 26
    assert [ws.cell(3, c).value for c in (TOT, TOT+1, TOT+2)] == ['Reçus', 'Envoyés', 'Taux']
    for r in range(4, 44):
        ref = 'Enfants!A%d' % (r + 3)
        ws.cell(r, TOT).value   = '=IF(%s="","",COUNTIF(D%d:Y%d,"O"))' % (ref, r, r)
        ws.cell(r, TOT+1).value = '=IF(%s="","",COUNTA(D%d:Y%d))' % (ref, r, r)
        ws.cell(r, TOT+2).value = '=IF(%s="","",IF(AA%d=0,"",Z%d/AA%d))' % (ref, r, r, r)
    fin = get_column_letter(TOT + 2)
    ws.merge_cells('A1:%s1' % fin); ws.merge_cells('A2:%s2' % fin)
    ws['A1'] = 'RÉVISIONS À LA MAISON — UN DEVOIR PAR SEMAINE'
    ws['A2'] = ("Un devoir envoyé aux parents chaque mercredi, à faire avant le samedi.  "
                "O = l'enfant a répondu · N = rien reçu · case vide = devoir pas encore échu.  "
                "Les noms et les totaux se remplissent tout seuls : ne tapez que dans les cases jaunes.")
    for i in range(22):
        ws.column_dimensions[get_column_letter(4 + i)].width = 5.6
    for c in (TOT, TOT+1, TOT+2):
        ws.column_dimensions[get_column_letter(c)].width = 9.0
    print('Révisions maison : restructure en 22 colonnes')

O, N = 'O', 'N'
ETAT = {
 'Ephraïm Christ-Emmanuel AMON': (N, N, None), 'Gédéon OKITONGA': (O, O, None),
 'Joyce Bouam LALLE': (O, O, O),               'Joseph Daniel NGAPELA': (N, N, None),
 'Lohiss LOUA': (O, O, O),                     'Caleb DOGBRE': (O, N, None),
 'Grâce-Elsa DREESEN': (O, O, None),           'Yanis AMESSAN': (O, O, O),
 'Isaac YEDOH LOHOUESS': (N, O, O),            'Acquilas MBIERE': (N, O, O),
 'Bénicia Chryti Léa MBIERE': (N, O, O),       'Viesainte MBIERE': (N, O, O),
 'Joshua DOGBRE': (O, N, None),                'Sarah LOUA': (O, O, None),
 'Ethan SORO': (N, N, None),                   'Els MOADJIDIBAYE': (O, O, None),
 'Eunice SORO': (N, N, None),
}
enf = wb['Enfants']; poses = 0
for r in range(4, 44):
    nom = enf.cell(r + 3, 1).value
    if not nom: continue
    nom = str(nom).strip()
    if nom not in ETAT: raise SystemExit('nom inconnu : %r' % nom)
    for i, v in enumerate(ETAT[nom]):
        if v is not None and ws.cell(r, 4 + i).value in (None, ''):
            ws.cell(r, 4 + i).value = v; poses += 1
print('Révisions maison :', poses, 'cases de devoir remplies')

# =====================================================================
#  2. OBSERVATIONS : la trace des retards excuses
# =====================================================================
obs = wb['Observations']
deja = {str(obs.cell(r, 5).value).strip() for r in range(4, 60)
        if obs.cell(r, 6).value == 'Ponctualité'}
ligne = 4
while obs.cell(ligne, 2).value: ligne += 1
ajouts = 0
for nom in ('Acquilas MBIERE', 'Bénicia Chryti Léa MBIERE',
            'Viesainte MBIERE', 'Isaac YEDOH LOHOUESS'):
    if nom in deja: continue
    r = ligne + ajouts
    obs.cell(r, 2).value = datetime.date(2026, 10, 2)
    obs.cell(r, 3).value = 'S01-S03'
    obs.cell(r, 4).value = 'JEREMIAH GEEKS'
    obs.cell(r, 5).value = nom
    obs.cell(r, 6).value = 'Ponctualité'
    obs.cell(r, 7).value = ("Arrive après 12 h : cours à son école le samedi matin. "
                            "Les retards des séances 1 à 3 sont excusés.")
    obs.cell(r, 8).value = ("Ne pas relancer la famille. La séance commence par la prière : "
                            "l'essentiel de la leçon est préservé.")
    obs.cell(r, 9).value = 'Permanent'
    obs.cell(r, 10).value = 'Oui'
    ajouts += 1
print('Observations :', ajouts, 'lignes ajoutées')

wb.save(DST)
print('classeur ecrit :', DST)
